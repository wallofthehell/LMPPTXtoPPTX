#!/usr/bin/env python3
"""NotebookLM/Gemini PPTX(슬라이드=통이미지) -> 텍스트 편집 가능 PPTX.
1단계(인식): python lmpptx2pptx.py analyze in.pptx draft.json [--tessdata DIR]   -> draft.json + 슬라이드별 미리보기 PNG
   draft.json 에서 text 를 고치고, 필요없는 줄은 "skip": true 로 바꾼다 (OCR 은 틀릴 수 있음: 'AI로' -> '시로' 등)
2단계(변환): python lmpptx2pptx.py build in.pptx draft.json out.pptx [--font "맑은 고딕"]
필요: pip install python-pptx opencv-python-headless pytesseract numpy pillow  + Tesseract(kor)"""
import argparse, io, re
import cv2, numpy as np, pytesseract
from PIL import Image
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN

def ocr_lines(img, lang, cfg):
    big = img
    words = []
    for k, im in enumerate((big, 255 - big)):          # 일반 + 반전(밝은 글자/어두운 배경)
        d = pytesseract.image_to_data(im, lang=lang, config=cfg + ' --psm 11', output_type=pytesseract.Output.DICT)
        for i, t in enumerate(d['text']):
            if t.strip() and float(d['conf'][i]) >= 75:
                x, y, w, h = [d[a][i] for a in ('left', 'top', 'width', 'height')]
                if k and any(min(x+w, a[2])-max(x, a[0]) > .3*w and min(y+h, a[3])-max(y, a[1]) > .3*h for a in (q['b'] for q in words if q['k'] == 0)):
                    continue                              # 반전 결과가 일반 결과와 겹치면 버림
                words.append(dict(k=k, key=(k, d['block_num'][i], d['par_num'][i], d['line_num'][i]), t=t, b=(x, y, x+w, y+h)))
    lines = {}
    for w in words: lines.setdefault(w['key'], []).append(w)
    out = []
    for ws in lines.values():
        ws.sort(key=lambda w: w['b'][0]); txt = ' '.join(w['t'] for w in ws)
        mh = np.median([w['b'][3]-w['b'][1] for w in ws]); cy = np.median([(w['b'][1]+w['b'][3])/2 for w in ws])
        b = [min(w['b'][0] for w in ws), cy-mh/2, max(w['b'][2] for w in ws), cy+mh/2]
        if 'Notebook' not in txt and re.search(r'[가-힣A-Za-z0-9]{2,}', txt) and 10 < b[3]-b[1] < 160 and (b[2]-b[0]) > 20:
            out.append(dict(t=txt, b=[int(v) for v in b]))
    return sorted(out, key=lambda l: (l['b'][1], l['b'][0]))

def group(lines):                                         # 인접 줄 -> 문단(텍스트 상자)
    boxes = []
    for l in lines:
        h = l['b'][3]-l['b'][1]; cx = (l['b'][0]+l['b'][2])/2
        for g in boxes:
            p = g[-1]; ph = p['b'][3]-p['b'][1]
            if 0 < l['b'][1]-p['b'][1] < 1.9*ph and abs(h-ph) < .3*ph and (abs(l['b'][0]-p['b'][0]) < .6*h or abs(cx-(p['b'][0]+p['b'][2])/2) < .6*h):
                g.append(l); break
        else: boxes.append([l])
    return boxes

def style(img, b, pad=3):
    H, W = img.shape[:2]; x0, y0, x1, y1 = max(b[0]-pad, 0), max(b[1]-pad, 0), min(b[2]+pad, W), min(b[3]+pad, H)
    ring = np.concatenate([img[max(y0-6, 0):y0, x0:x1].reshape(-1, 3), img[y1:y1+6, x0:x1].reshape(-1, 3),
                           img[y0:y1, max(x0-6, 0):x0].reshape(-1, 3), img[y0:y1, x1:x1+6].reshape(-1, 3)])
    bg = np.median(ring, 0) if len(ring) else np.array([255, 255, 255])
    sub = img[y0:y1, x0:x1].astype(float); dist = np.linalg.norm(sub - bg, axis=2); m = dist > 60
    if m.sum() < 10: return None
    far = dist[m] >= np.percentile(dist[m], 80)
    col = img[y0:y1, x0:x1][m][far].mean(0)[::-1].astype(int)
    return (x0, y0, x1, y1), m, tuple(int(c) for c in col), m.sum()/((x1-x0)*(y1-y0))

def pics_of(prs, sl):
    sw, sh = prs.slide_width, prs.slide_height
    return [s for s in sl.shapes if s.shape_type == 13 and s.width >= .9*sw and s.height >= .9*sh]

def load(part):
    return cv2.cvtColor(np.array(Image.open(io.BytesIO(part.blob)).convert('RGB')), cv2.COLOR_RGB2BGR)

def analyze(src, dst, lang='kor+eng', tessdata=None):
    cfg = f'--tessdata-dir {tessdata}' if tessdata else ''; prs = Presentation(src); res = []
    for n, sl in enumerate(prs.slides, 1):
        for pic in pics_of(prs, sl):
            img = load(sl.part.related_part(pic._element.blipFill.blip.rEmbed)); prev = img.copy()
            for gi, g in enumerate(group(ocr_lines(img, lang, cfg))):
                for l in g:
                    res.append(dict(slide=n, pic=pic.shape_id, group=gi, text=l['t'], box=l['b'], skip=False))
                    cv2.rectangle(prev, tuple(l['b'][:2]), tuple(l['b'][2:]), (0, 0, 255), 2)
                    cv2.putText(prev, str(len(res)-1), (l['b'][0], max(l['b'][1]-4, 12)), 0, .6, (0, 0, 255), 2)
            cv2.imwrite(dst.rsplit('.', 1)[0] + f'_slide{n}.png', prev)
    import json; json.dump(res, open(dst, 'w', encoding='utf-8'), ensure_ascii=False, indent=1); return len(res)

def build(src, draft, dst, font='맑은 고딕'):
    import json; items = [r for r in json.load(open(draft, encoding='utf-8')) if not r.get('skip') and r['text'].strip()]
    prs = Presentation(src); stats = []
    for n, sl in enumerate(prs.slides, 1):
        for pic in pics_of(prs, sl):
            mine = [r for r in items if r['slide'] == n and r['pic'] == pic.shape_id]
            if not mine: continue
            part = sl.part.related_part(pic._element.blipFill.blip.rEmbed); img = load(part)
            Hh, Ww = img.shape[:2]; k = pic.width/Ww; mask = np.zeros((Hh, Ww), np.uint8); boxes = {}
            for r in mine:
                s = style(img, r['box'])
                if not s: continue
                bb, m, col, ink = s; mask[bb[1]:bb[3], bb[0]:bb[2]] |= m.astype(np.uint8)*255
                boxes.setdefault(r['group'], []).append((dict(t=r['text'], b=r['box']), col, ink))
            out = cv2.inpaint(img, cv2.dilate(mask, np.ones((5, 5), np.uint8)), 5, cv2.INPAINT_TELEA)
            buf = io.BytesIO(); Image.fromarray(cv2.cvtColor(out, cv2.COLOR_BGR2RGB)).save(buf, 'PNG'); part._blob = buf.getvalue()
            pic.name = '배경(원본 디자인)'
            for i, rows in enumerate(boxes.values(), 1):
                ls = [l for l, _, _ in rows]; col = rows[0][1]; ink = np.mean([x[2] for x in rows])
                x0 = min(l['b'][0] for l in ls); x1 = max(l['b'][2] for l in ls); y0 = ls[0]['b'][1]; y1 = ls[-1]['b'][3]
                lh = np.median([l['b'][3]-l['b'][1] for l in ls])
                centered = len(ls) > 1 and max(abs((l['b'][0]+l['b'][2])/2-(x0+x1)/2) for l in ls) < .6*lh and max(abs(l['b'][0]-x0) for l in ls) > .8*lh
                tb = sl.shapes.add_textbox(pic.left+Emu(int(x0*k)), pic.top+Emu(int(y0*k)), Emu(int((x1-x0)*k)), Emu(int((y1-y0)*k)))
                tb.name = f'텍스트{i}'; tf = tb.text_frame; tf.word_wrap = False; tf.vertical_anchor = MSO_ANCHOR.MIDDLE
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                pt = lh*k/12700*.93; bold = ink > .2
                for j, l in enumerate(ls):
                    p = tf.paragraphs[0] if j == 0 else tf.add_paragraph(); p.alignment = PP_ALIGN.CENTER if centered else PP_ALIGN.LEFT
                    if len(ls) > 1: p.line_spacing = Pt((y1-y0)/(len(ls)-.15)*k/12700)
                    r = p.add_run(); r.text = l['t']; f = r.font; f.size = Pt(round(pt, 1)); f.bold = bold; f.name = font; f.color.rgb = RGBColor(*col)
            stats.append((n, len(boxes)))
    prs.save(dst); return stats

if __name__ == '__main__':
    a = argparse.ArgumentParser(); a.add_argument('cmd', choices=['analyze', 'build']); a.add_argument('files', nargs='+')
    a.add_argument('--lang', default='kor+eng'); a.add_argument('--font', default='맑은 고딕'); a.add_argument('--tessdata'); a = a.parse_args()
    if a.cmd == 'analyze': print(f'{analyze(*a.files[:2], a.lang, a.tessdata)}개 줄 인식 -> 미리보기 PNG 와 JSON 을 확인하세요')
    else:
        for n, c in build(*a.files[:3], a.font): print(f'슬라이드 {n}: 텍스트 상자 {c}개')
